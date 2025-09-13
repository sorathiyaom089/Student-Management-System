# SPDX-License-Identifier: Apache-2.0
"""
This example shows how to use vLLM for running offline inference with
the correct prompt format on vision language models for text generation.

For most models, the prompt format should follow corresponding examples
on HuggingFace model repository.
"""
import os
import random
from PIL import Image
import fitz
from docx import Document
from huggingface_hub import snapshot_download
from transformers import AutoTokenizer

# The vllm library is required. You can install it with `pip install vllm`.
from vllm import LLM, SamplingParams
from vllm.assets.image import ImageAsset
from vllm.assets.video import VideoAsset
from vllm.lora.request import LoRARequest
from vllm.utils import FlexibleArgumentParser


os.environ["no_proxy"] = "localhost,127.0.0.1,::1"
import gradio as gr

# NOTE: The model paths below are placeholders.
# You must replace these with the actual paths to your downloaded models
# or the model names from HuggingFace Hub.
# The code cannot be run without a valid model.
try:
    mllm = LLM(
        model="OceanGPT-V's path",
        max_model_len=4096,
        max_num_seqs=5,
        mm_processor_kwargs={
            "min_pixels": 28 * 28,
            "max_pixels": 1280 * 28 * 28,
            "fps": 1,
        }
    )

    llm = LLM(model="OceanGPT's path")

    coder = LLM(model="OceanGPT-coder's path")
except Exception as e:
    print(f"Error initializing LLM models. Please check your model paths. Error: {e}")
    mllm = None
    llm = None
    coder = None

def extract_text_from_file(file_path):
    """
    Extracts text from a PDF or DOCX file.
    """
    if file_path and file_path.lower().endswith(".pdf"):
        text = ""
        with fitz.open(file_path) as doc:
            for page in doc:
                text += page.get_text()
        return text
    elif file_path and file_path.lower().endswith(".docx"):
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    else:
        return ""
    
# Qwen2.5
def chat_qwen(questions: list[str], llm_file, temperature: float, top_p: float, max_tokens: int):
    """
    Generates a response using the text-based LLM.
    """
    if llm is None:
        return "Model not initialized. Please check the console for errors."

    sampling_params = SamplingParams(temperature=temperature, top_p=top_p, max_tokens=max_tokens)
    if llm_file:
        text = extract_text_from_file(llm_file)
        # Prepend extracted text to the question
        questions = [text + '\n' + question for question in questions]
    outputs = llm.generate(questions, sampling_params)[0]
    generated_text = outputs.outputs[0].text
    return generated_text

def chat_qwen_coder(questions: list[str], temperature: float, top_p: float, max_tokens: int):
    """
    Generates a response using the coder LLM.
    """
    if coder is None:
        return "Model not initialized. Please check the console for errors."

    sampling_params = SamplingParams(temperature=temperature, top_p=top_p, max_tokens=max_tokens)
    outputs = coder.generate(questions, sampling_params)[0]
    generated_text = outputs.outputs[0].text
    return generated_text
 
# Qwen2.5-VL
def run_qwen2_5_vl(questions: list[str], modality: str):
    """
    Formats the prompts for the multimodal model.
    """
    if modality == "image":
        placeholder = "<|image_pad|>"
    elif modality == "video":
        placeholder = "<|video_pad|>"

    prompts = [
        ("<|im_start|>system\nYou are a helpful assistant.<|im_end|>\n"
         f"<|im_start|>user\n<|vision_start|>{placeholder}<|vision_end|>"
         f"{question}<|im_end|>\n"
         "<|im_start|>assistant\n") for question in questions
    ]
    stop_token_ids = None
    return prompts, stop_token_ids


def get_multi_modal_input(img_questions, image_path):
    """
    Prepares the multimodal input for the LLM.
    """
    if not image_path:
        return {"data": None, "questions": img_questions}

    # Input image and question
    image = Image.open(image_path).convert("RGB")

    return {
        "data": image,
        "questions": img_questions,
    }


def chat_with_qwenvl(img_question: str, image_path: str, temperature: float, top_p: float, max_tokens: int):
    """
    Generates a response using the multimodal LLM.
    """
    if mllm is None:
        return "Model not initialized. Please check the console for errors."

    modality = "image"
    
    # Wrap the single question string in a list for the function call
    questions = [img_question]
    mm_input = get_multi_modal_input(questions, image_path)
    data = mm_input["data"]
    
    if not data:
        return "No image provided."

    prompts, stop_token_ids = run_qwen2_5_vl(questions, modality)
    
    sampling_params = SamplingParams(temperature=temperature, top_p=top_p, max_tokens=max_tokens,
                                     stop_token_ids=stop_token_ids)

    inputs = {
        "prompt": prompts[0],
        "multi_modal_data": {
            modality: data
        },
    }

    outputs = mllm.generate(inputs, sampling_params=sampling_params)

    return outputs[0].outputs[0].text

def check_file_size(file):
    """
    Checks the size of the uploaded file.
    """
    if file is None:
        return gr.update(visible=False)
    
    if isinstance(file, str):
        file_path = file
    elif isinstance(file, dict) and "name" in file:
        file_path = file.get("name")
    else:
        return gr.update(visible=False)
        
    size_in_bytes = os.path.getsize(file_path)
    if size_in_bytes > 1 * 1024 * 1024:
        return gr.update(value="⚠️ Uploaded file exceeds 1MB, please upload a smaller file.", visible=True)
    else:
        return gr.update(visible=False)

def create_demo():
    """
    Creates the Gradio UI.
    """
    with gr.Blocks(css="""
        .textarea-auto-wrap textarea {
            white-space: pre-wrap !important;
            word-wrap: break-word !important;
            overflow-x: hidden !important;
            overflow-y: auto !important;
            resize: vertical !important;
        }
        .textarea-fixed-height textarea {
            white-space: pre-wrap !important;
            word-wrap: break-word !important;
            overflow-x: hidden !important;
            overflow-y: auto !important;
            resize: none !important;
            max-height: 600px !important;
        }
        .image-with-scroll {
            max-height: 400px !important;
            overflow: auto !important;
        }
        .image-with-scroll img {
            max-width: 100% !important;
            height: auto !important;
        }
    """) as demo:
        with gr.Tab("OceanGPT-o"):    
            with gr.Row():
                with gr.Column():
                    mllm_text = gr.TextArea(
                        placeholder="Input text query", 
                        label="text input", 
                        lines=3,
                        max_lines=15,
                        elem_classes=["textarea-auto-wrap"]
                    )
                    file_warning = gr.Markdown("", visible=False)
                    mllm_image = gr.Image(type="filepath", label="image input", elem_classes=["image-with-scroll"])
                    temperature = gr.Slider(minimum=0, maximum=2, label="temperature", step=0.1, value=0.6)
                    top_p = gr.Slider(minimum=0, maximum=1, label="top_p", step=0.01, value=1.0)
                    max_tokens = gr.Slider(minimum=1, maximum=4096, label="max_tokens", step=1, value=2048)
                    clear_button = gr.ClearButton(components=[mllm_text, mllm_image],value="Clear")
                    run_button = gr.Button("Run")
                with gr.Column():
                    response_res = gr.TextArea(
                        label="OceanGPT-o's response", 
                        lines=8,
                        max_lines=20,
                        elem_classes=["textarea-fixed-height"]
                    )
            
            mllm_image.change(fn=check_file_size, inputs=mllm_image, outputs=[file_warning])
            
            inputs = [mllm_text, mllm_image, temperature, top_p, max_tokens]
            outputs = [response_res]
            
            # NOTE: Removed gr.Examples because the referenced files are not available.
            
            clear_button.add([response_res])
            run_button.click(fn=chat_with_qwenvl,
                             inputs=inputs, outputs=outputs)
            
        with gr.Tab("OceanGPT"):
            with gr.Row():
                with gr.Column():
                    llm_text = gr.TextArea(
                        placeholder="Input query", 
                        label="text input", 
                        lines=3,
                        max_lines=15,
                        elem_classes=["textarea-auto-wrap"]
                    )
                    file_warning = gr.Markdown("", visible=False)
                    llm_file = gr.File(label="Upload PDF / Word", file_types=[".pdf", ".docx"])
                    temperature = gr.Slider(minimum=0, maximum=2, label="temperature", step=0.1, value=0.6)
                    top_p = gr.Slider(minimum=0, maximum=1, label="top_p", step=0.01, value=1.0)
                    max_tokens = gr.Slider(minimum=1, maximum=4096, label="max_tokens", step=1, value=2048)
                    clear_button = gr.ClearButton(components=[llm_text, llm_file],value="Clear")
                    run_button = gr.Button("Run")
                with gr.Column():
                    llm_response_res = gr.TextArea(
                        label="OceanGPT's response", 
                        lines=8,
                        max_lines=20,
                        elem_classes=["textarea-fixed-height"]
                    )
            
            llm_file.change(fn=check_file_size, inputs=llm_file, outputs=[file_warning])

            inputs = [llm_text, llm_file, temperature, top_p, max_tokens]
            outputs = [llm_response_res]
            
            # NOTE: Removed gr.Examples because the referenced files are not available.
            
            clear_button.add([llm_response_res])
            # Pass a list containing the single text input to the chat function
            run_button.click(fn=lambda text, file, temp, top_p, max_t: chat_qwen([text], file, temp, top_p, max_t),
                             inputs=inputs, outputs=outputs)
            
        with gr.Tab("OceanGPT-coder"):
            with gr.Row():
                with gr.Column():
                    llm_text = gr.TextArea(
                        placeholder="Input query", 
                        label="text input", 
                        lines=3,
                        max_lines=15,
                        elem_classes=["textarea-auto-wrap"]
                    )
                    temperature = gr.Slider(minimum=0, maximum=2, label="temperature", step=0.1, value=0.6)
                    top_p = gr.Slider(minimum=0, maximum=1, label="top_p", step=0.01, value=1.0)
                    max_tokens = gr.Slider(minimum=1, maximum=4096, label="max_tokens", step=1, value=2048)

                    clear_button = gr.ClearButton(components=[llm_text],value="Clear")
                    run_button = gr.Button("Run")
                with gr.Column():
                    llm_response_res = gr.TextArea(
                        label="OceanGPT-coder's response", 
                        lines=8,
                        max_lines=20,
                        elem_classes=["textarea-fixed-height"]
                    )
            
            inputs = [llm_text, temperature, top_p, max_tokens]
            outputs = [llm_response_res]
            
            # NOTE: Removed gr.Examples because the referenced files are not available.
            
            clear_button.add([llm_response_res])
            # Pass a list containing the single text input to the chat function
            run_button.click(fn=lambda text, temp, top_p, max_t: chat_qwen_coder([text], temp, top_p, max_t),
                             inputs=inputs, outputs=outputs)
            
        with gr.Accordion("Limitations"):
            gr.Markdown("""
            - The model may have hallucination issues.
            - Due to limited computational resources, OceanGPT-o currently only supports natural language generation for certain types of sonar images and ocean science images. OceanGPT-coder currently only supports MOOS code generation.
            - We did not optimize the identity and the model may generate identity information similar to that of Qwen/MiniCPM/LLaMA/GPT series models.
            - The model's output is influenced by prompt tokens, which may result in inconsistent results across multiple attempts.
            """)
    return demo

description = """
# Ocean Foundation Models
Upload documents (Word, PDF, or images) to help OceanGPT provide more accurate answers.

Please refer to our [project](http://www.oceangpt.blue/) for more details.
"""

# The main block to launch the Gradio app
with gr.Blocks(css="""
    h1,p {text-align: center !important;}
    .textarea-auto-wrap textarea {
        white-space: pre-wrap !important;
        word-wrap: break-word !important;
        overflow-x: hidden !important;
        overflow-y: auto !important;
        resize: vertical !important;
    }
    .textarea-fixed-height textarea {
        white-space: pre-wrap !important;
        word-wrap: break-word !important;
        overflow-x: hidden !important;
        overflow-y: auto !important;
        resize: none !important;
        max-height: 600px !important;
    }
    .image-with-scroll {
        max-height: 400px !important;
        overflow: auto !important;
    }
    .image-with-scroll img {
        max-width: 100% !important;
        height: auto !important;
    }
""") as demo:
    gr.Markdown(description)
    create_demo()

demo.queue().launch(server_name="0.0.0.0",server_port=7860)
